# -*- coding: utf-8 -*-
"""Kampanya Takip — ekran (Ürün Yönetimi › Kampanya Takip).

YAPI
  Üst özet      : sürüyor · kapanmayı bekliyor · bu yıl destek · bu yıl net kâr
  Uyarı         : süresi dolmuş ama kapatılmamış kampanyalar (adet girilmeli)
  Komut çubuğu  : görünüm (Güncel / Kapanmayı bekleyen / Kapalı / Tümü) · arama
                  · filtre (müşteri, kategori, yıl) · Yeni kampanya
  Kartlar       : zaman çizgisi + ürün/adet + destek + net kâr + marj;
                  kartın TAMAMI tıklanır → detay penceresi
  Detay         : Ürünler (tabloda düzenle, ekle, sil, kapat) · Bilgiler · İşlemler

Rakamların hepsi kampanya_hesap.py'den gelir (tek formül, testli).
Excel şablonu akışı (kampanya + ürünler tek dosyada) Ekim 2026'dan beri üst menüdeki
Dosya kapısında: kapi_kampanya (şablon: sablon_kampanya).
"""
import html as _h
from io import BytesIO

import pandas as pd
import streamlit as st

from shared.tasarim import (baslik, css_tek_satir, kpi_serit, mesaj, bos_durum,
                            rv, sayi, tr_sayi)
from shared.utils import firma_gorunen_ad, normalize_tr, tr_kucuk, tr_today
from shared.ana_veri import kategori_ad as _kat_ad, kategori_anahtar as _kat_anh, urun_ad as _urun_ad   # tek kaynak (Eki 2026)
from shared import bilesen as B
from . import kampanya_hesap as H
from .analitik import tum_urunler_listesi
from .database import (_cache_temizle, ekle_kampanya, ekle_kampanya_urun, get_client,
                       get_kampanyalar, get_kampanya_urunler, get_tum_kampanya_urunler,
                       guncelle_kampanya, guncelle_kampanya_urun, kapat_kampanya,
                       sil_kampanya, sil_kampanya_urun)

_AY = ["", "Oca", "Şub", "Mar", "Nis", "May", "Haz", "Tem", "Ağu", "Eyl", "Eki", "Kas", "Ara"]
_GORUNUM = {"guncel": "Güncel", "bekliyor": "Kapanmayı bekleyen", "kapali": "Kapalı", "tumu": "Tümü"}


# ════════════════════════════════════════════════════════════════════
# Yardımcılar
# ════════════════════════════════════════════════════════════════════
def _usd(v, kisa=True):
    if v is None:
        return "—"
    return sayi(v, "$") if kisa else f"{'-' if v < 0 else ''}${tr_sayi(abs(v), 2)}"


def _tarih_kisa(s):
    try:
        y, m, d = str(s)[:10].split("-")
        return f"{int(d)} {_AY[int(m)]}"
    except (ValueError, IndexError):
        return "—"


def _aralik(k):
    b, e = str(k.get("baslangic_tarihi") or ""), str(k.get("bitis_tarihi") or "")
    yil = e[:4] if e[:4] == b[:4] else ""
    return f"{_tarih_kisa(b)} – {_tarih_kisa(e)}{(' ' + yil) if yil else (' ' + e[:4] if e else '')}"


def _kalan_metni(o):
    d, kalan = o["durum"], o["kalan"]
    if d == "kapali":
        return "Kapandı"
    if kalan is None:
        return "Tarih yok"
    if d == "yaklasan":
        bas = o["gun"] - kalan - 1
        return f"{-bas} gün sonra başlıyor" if bas < 0 else "Yarın başlıyor"
    if d == "bekliyor":
        return f"{-kalan} gün önce bitti"
    return "Bugün bitiyor" if kalan == 0 else f"{kalan} gün kaldı"


def _firma_ad(kod):
    try:
        return firma_gorunen_ad(kod, kisa=False) or str(kod or "")   # cari kartındaki tam ad
    except Exception:  # noqa: BLE001 — eşleme tablosu yoksa kodu göster
        return str(kod or "")


def _firma_secenekleri(kamps=()):
    """Firma seçenekleri VERİDEN (KANAL yok, her firma kendi adıyla — Ekim 2026)."""
    try:
        from .database import get_firma_listesi
        veri = list(get_firma_listesi() or [])
    except Exception:  # noqa: BLE001
        veri = []
    return H.firma_secenekleri(veri + [k.get("firma") for k in (kamps or ())])


def _kategoriler(urunler, kamps):
    """Kategori seçenekleri: kayıt değeri (küçük harf) aynı kalır; anahtara göre TEKİL, tek yazım."""
    _kv = {}
    for _v in [u.get("kategori") for u in urunler] + [k.get("kategori") for k in kamps]:
        if tr_kucuk(_v):
            _kv.setdefault(_kat_anh(_v), tr_kucuk(_v))
    return sorted(_kv.values(), key=lambda x: _kat_ad(x).lower())


def _veri():
    """Kampanyalar, ürünleri ve güncel paçallar (hepsi önbellekli kaynaklardan)."""
    kamps = get_kampanyalar() or []
    ku_map = {}
    for r in get_tum_kampanya_urunler() or []:
        ku_map.setdefault(r.get("kampanya_id"), []).append(r)
    try:
        urunler = tum_urunler_listesi() or []
    except Exception as e:  # noqa: BLE001
        from shared.hata_log import kaydet
        kaydet("kampanya.urunler", e)
        urunler = []
    pacal = {u["sku"]: float(u.get("final_cost_price") or 0) for u in urunler}
    return kamps, ku_map, pacal, urunler


def _yeniden_ac(kid):
    get_client().table("kampanyalar").update({"durum": "aktif"}).eq("id", kid).execute()
    _cache_temizle()


def _detay_ac(kid):
    B.detay_ac("kmp", kid)


# ════════════════════════════════════════════════════════════════════
# Stil
# ════════════════════════════════════════════════════════════════════
def _css():
    return "<style>" + css_tek_satir("""
.kmp-ust{display:flex;align-items:flex-start;gap:10px;}
.kmp-ad{flex:1;min-width:0;font-size:14.5px;font-weight:650;color:var(--k-metin);line-height:1.3;
  overflow:hidden;text-overflow:ellipsis;white-space:nowrap;}
.kmp-meta{margin-top:3px;font-size:12px;color:var(--k-soluk);white-space:nowrap;overflow:hidden;text-overflow:ellipsis;}
.kmp-zaman{display:flex;align-items:center;gap:10px;margin:12px 0 12px;font-size:12px;color:var(--k-silik);}
.kmp-cubuk{flex:1;height:5px;border-radius:99px;background:var(--k-ortu2);overflow:hidden;}
.kmp-cubuk i{display:block;height:100%;border-radius:99px;background:var(--d);}
.kmp-zaman b{font-weight:600;color:var(--k-soluk);white-space:nowrap;}
.kmp-zaman span{white-space:nowrap;}
.kmp-rakam{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:8px;padding-top:11px;
  border-top:1px solid var(--k-kenar);}
.kmp-rakam div{min-width:0;}
.kmp-rakam small{display:block;font-size:11.5px;color:var(--k-silik);white-space:nowrap;}
.kmp-rakam b{display:block;margin-top:1px;font-family:var(--k-mono);font-variant-numeric:tabular-nums;
  font-size:14px;font-weight:600;color:var(--k-metin);white-space:nowrap;overflow:hidden;text-overflow:ellipsis;}
.kmp-rakam b.poz{color:var(--k-yesil);} .kmp-rakam b.neg{color:var(--k-kirmizi);} .kmp-rakam b.yok{color:var(--k-silik);font-family:inherit;font-weight:400;font-size:12.5px;}
.kmp-dt{display:flex;flex-direction:column;gap:8px;margin:-4px 0 14px;}
.kmp-dt-ust{display:flex;align-items:center;gap:8px;flex-wrap:wrap;}
.kmp-dt-ust .kmp-ad{flex:0 1 auto;font-size:19px;white-space:normal;}
.kmp-dt .kmp-zaman{margin:2px 0 0;}
.kmp-bos{font-size:12.5px;color:var(--k-soluk);margin:2px 0 10px;}
[data-testid="stMarkdownContainer"]:has(> .k-mesaj){margin-bottom:6px !important;}
[data-testid="stHorizontalBlock"]:has([class*="st-key-tk_kart_kmp"]){margin-bottom:2px;}
@media (max-width:640px){ .kmp-rakam{grid-template-columns:repeat(2,minmax(0,1fr));} }
""") + "</style>"


# ════════════════════════════════════════════════════════════════════
# Kart
# ════════════════════════════════════════════════════════════════════
def _kart_html(k, o):
    ad_d, renk, _ = H.DURUMLAR[o["durum"]]
    meta = " · ".join(x for x in (_firma_ad(k.get("firma")), _kat_ad(k.get("kategori")),
                                   (k.get("kampanya_turu") or "").strip()) if x)
    if o["adet"] > 0:
        net_sinif = "poz" if o["net"] >= 0 else "neg"
        net = f'<b class="{net_sinif}">{_usd(o["net"])}</b>'
        marj = f'<b>%{tr_sayi(o["marj"], 1)}</b>' if o["marj"] is not None else '<b class="yok">—</b>'
        destek = f'<b>{_usd(o["destek"])}</b>'
    else:
        _m = "adet girilmedi" if o["durum"] in ("bekliyor", "kapali") else "satış bekleniyor"
        destek = '<b class="yok">—</b>'
        net = f'<b class="yok">{_m}</b>'
        marj = '<b class="yok">—</b>'
    oran = 1.0 if o["durum"] == "kapali" else o["oran"]
    return (f'<div style="--d:{rv(renk)}">'
            f'<div class="kmp-ust"><div class="kmp-ad" title="{_h.escape(str(k.get("kampanya_adi","")), quote=True)}">'
            f'{_h.escape(str(k.get("kampanya_adi") or "—"))}</div>{B.cip(ad_d, renk)}</div>'
            f'<div class="kmp-meta">{_h.escape(meta) or "&nbsp;"}</div>'
            f'<div class="kmp-zaman"><span>{_aralik(k)}</span><div class="kmp-cubuk"><i style="width:{oran*100:.0f}%"></i></div>'
            f'<b>{_kalan_metni(o)}</b></div>'
            f'<div class="kmp-rakam">'
            f'<div><small>{o["urun"]} ürün</small><b>{tr_sayi(o["adet"])} adet</b></div>'
            f'<div><small>Destek</small>{destek}</div>'
            f'<div><small>Net kâr</small>{net}</div>'
            f'<div><small>Net marj</small>{marj}</div>'
            f'</div></div>')


# ════════════════════════════════════════════════════════════════════
# Ana ekran
# ════════════════════════════════════════════════════════════════════
def render():
    st.markdown(_css(), unsafe_allow_html=True)
    bugun = tr_today()
    kamps, ku_map, pacal, urunler = _veri()
    _katlar = _kategoriler(urunler, kamps)
    # Başlık + ana eylem aynı satırda
    h1, h3 = st.columns([6.45, 1.55], vertical_alignment="center")
    h1.markdown(baslik("🎯 Ürün Yönetimi", "Kampanya Takip",
                       aciklama="Firma destekli kampanyalar: süre, kârlılık ve kapanış tek ekranda."),
                unsafe_allow_html=True)
    if h3.button("Yeni kampanya", type="primary", icon=":material/add:", use_container_width=True,
                 key="kmp_yeni_btn"):
        _yeni_dialog(_katlar)
    oz = {k["id"]: H.kampanya_ozet(k, ku_map.get(k["id"], []), pacal, bugun) for k in kamps}
    say = {d: sum(1 for o in oz.values() if o["durum"] == d) for d in H.DURUM_SIRA}

    # ── Üst özet: bu yıl başlayan kampanyalar ──
    yil = str(bugun.year)
    bu_yil = [oz[k["id"]] for k in kamps if str(k.get("baslangic_tarihi") or "")[:4] == yil]
    _d = sum(o["destek"] for o in bu_yil)
    _fd = sum(o["firma_destek"] for o in bu_yil)
    _ed = sum(o["ek_destek"] for o in bu_yil)
    _net = sum(o["net"] for o in bu_yil if o["adet"] > 0)
    _ciro = sum(o["ciro"] for o in bu_yil if o["adet"] > 0)
    st.markdown(kpi_serit([
        {"etiket": "Sürüyor", "deger": tr_sayi(say["suruyor"]), "renk": "yesil",
         "alt": f"{say['yaklasan']} yaklaşan"},
        {"etiket": "Kapanmayı bekliyor", "deger": tr_sayi(say["bekliyor"]),
         "renk": "amber" if say["bekliyor"] else "silik",
         "alt": "satış adedi girilmeli" if say["bekliyor"] else "bekleyen yok"},
        {"etiket": f"{yil} destek", "deger": _usd(_d), "renk": "mor",
         "alt": f"firma {_usd(_fd)} · ek {_usd(_ed)}", "ipucu": f"${tr_sayi(_d, 2)}"},
        {"etiket": f"{yil} net kâr", "deger": _usd(_net), "renk": "yesil" if _net >= 0 else "kirmizi",
         "alt": (f"net marj %{tr_sayi(_net / _ciro * 100, 1)} · spiff sonrası" if _ciro > 0 else "spiff sonrası"),
         "ipucu": f"${tr_sayi(_net, 2)}"},
    ]), unsafe_allow_html=True)

    # ── Komut çubuğu ──
    _sec = [f"{_GORUNUM['guncel']} ({say['suruyor'] + say['yaklasan'] + say['bekliyor']})",
            f"{_GORUNUM['bekliyor']} ({say['bekliyor']})",
            f"{_GORUNUM['kapali']} ({say['kapali']})", f"Tümü ({len(kamps)})"]
    c1, c2, c3 = st.columns([4.2, 2.6, 1.1], vertical_alignment="bottom")
    with c1:
        _g = st.segmented_control("Görünüm", _sec, default=_sec[0], key="kmp_gorunum",
                                  label_visibility="collapsed") or _sec[0]
    gorunum = ["guncel", "bekliyor", "kapali", "tumu"][_sec.index(_g)] if _g in _sec else "guncel"
    ara = c2.text_input("Ara", key="kmp_ara", placeholder="Kampanya, firma, kategori…",
                        label_visibility="collapsed")
    _yillar = sorted({str(k.get("baslangic_tarihi") or "")[:4] for k in kamps
                      if str(k.get("baslangic_tarihi") or "")[:4].isdigit()}, reverse=True)
    _f = B.filtre(c3, [
        {"etiket": "Müşteri", "secenekler": _firma_secenekleri(kamps), "key": "kmp_f_firma",
         "format_func": lambda f: f if f in ("Tümü", "DİĞER") else _firma_ad(f)},
        {"etiket": "Kategori", "secenekler": _katlar, "key": "kmp_f_kat",
         "format_func": lambda x: x if x == "Tümü" else _kat_ad(x)},
        {"etiket": "Yıl", "secenekler": _yillar, "key": "kmp_f_yil"}])
    f_firma, f_kat, f_yil = _f["kmp_f_firma"], _f["kmp_f_kat"], _f["kmp_f_yil"]

    # ── Süresi dolmuşlar: işin kendisi, uyarıyla başla ──
    if say["bekliyor"] and gorunum in ("guncel", "tumu"):
        st.markdown(mesaj("uyari", f"{say['bekliyor']} kampanyanın süresi doldu ama kapatılmadı. "
                                   "Karta tıkla, Ürünler sekmesinde satış adetlerini gir ve kapat — "
                                   "net kâr ancak o zaman kesinleşir."), unsafe_allow_html=True)

    _dsec = "tumu" if gorunum in ("guncel", "tumu") else gorunum
    liste = H.filtrele(kamps, bugun, _dsec, f_firma, f_kat, f_yil, ara)
    if gorunum == "guncel":
        liste = [k for k in liste if oz[k["id"]]["durum"] != "kapali"]

    if not liste:
        if not kamps:
            st.markdown(bos_durum("Henüz kampanya yok",
                                  "Sağ üstteki 'Yeni kampanya' ile başla ya da doldurulmuş Excel şablonunu üst "
                                  "menüdeki Dosya düğmesinden yükleyip kampanyayı ürünleriyle birlikte oluştur.", "campaign"),
                        unsafe_allow_html=True)
        else:
            st.markdown(bos_durum("Bu görünümde kampanya yok",
                                  "Filtreleri gevşet, aramayı temizle ya da başka bir görünüm seç.",
                                  "filter_alt_off"), unsafe_allow_html=True)
    else:
        _goster = liste[:int(st.session_state.get("kmp_limit", 24))]
        for i in range(0, len(_goster), 2):
            cols = st.columns(2, gap="small")
            for col, k in zip(cols, _goster[i:i + 2]):
                o = oz[k["id"]]
                with col:
                    B.tiklanir(f"kmp{k['id']}", _kart_html(k, o), _detay_ac, (k["id"],), tur="kart",
                               renk=H.DURUMLAR[o["durum"]][1], etiket=f"{k.get('kampanya_adi','')} detayını aç")
        if len(liste) > len(_goster):
            if st.button(f"Daha fazla göster ({len(liste) - len(_goster)} kampanya daha)",
                         key="kmp_daha", use_container_width=True, type="tertiary"):
                st.session_state["kmp_limit"] = int(st.session_state.get("kmp_limit", 24)) + 24
                st.rerun(scope="fragment")

    _sec = B.detay_istendi("kmp")
    if _sec:
        _detay_dialog(_sec)


# ════════════════════════════════════════════════════════════════════
# Detay penceresi
# ════════════════════════════════════════════════════════════════════
def _kapat_ve_ref_ac(kamp):
    """Kampanyayı kapatır; Ref No Takip'e kampanyanın desteğiyle ref açar (kayranpm.kampanya_ref) ve
    kapatan kişiye + İbrahim'e mail atar. Ref / mail tutmazsa kampanya yine kapanır. Ekran mesajı döner."""
    kid = kamp["id"]
    kapat_kampanya(kid)
    try:
        from kayranpm import kampanya_ref as KR
        from kayranpm import ref_no as N
        urunler = get_client().table("kampanya_urunler").select("*").eq("kampanya_id", kid).execute().data or []
        sonuc = KR.kapaninca_ref_ac(kamp, urunler, N.get_firmalar(), N.get_refler, N.ref_ekle_no, tr_today())
    except Exception as e:  # noqa: BLE001
        from shared.hata_log import kaydet
        kaydet("kampanya.kapanis_ref", e)
        return "Kampanya kapatıldı ama ref açılamadı (hata kaydı tutuldu); Ref No Takip'te elle gir."
    try:
        from shared import eposta as E
        alicilar = KR.mail_alicilari(st.session_state.get("aktif_kullanici", ""), E.adresler())
        if alicilar:
            konu, html = KR.mail_icerigi(kamp, sonuc, st.session_state.get("aktif_kullanici", ""))
            E.arka_planda(alicilar, konu, html)
    except Exception as e:  # noqa: BLE001 — mail gitmese de kapanış ve ref kayıtlı
        from shared.hata_log import kaydet
        kaydet("kampanya.kapanis_mail", e)
    return KR.ekran_mesaji(sonuc)


def _ref_onizleme(kamp):
    """Kapatmadan önce: hangi ref açılacak, aynı ay elle girilmiş ref var mı."""
    try:
        from kayranpm import kampanya_ref as KR
        from kayranpm import ref_no as N
        o = KR.onizleme(kamp, get_kampanya_urunler(kamp["id"]) or [], N.get_firmalar(), N.get_refler)
    except Exception:  # noqa: BLE001 — önizleme gösterilemese de kapatma çalışır
        return
    if o["mevcut"]:
        st.caption("Bu kampanyanın ref'i zaten var (" + ", ".join(r.get("ref_no") or "" for r in o["mevcut"])
                   + "); kapatınca yeni ref açılmaz.")
        return
    if o["sorun"]:
        st.markdown(mesaj("uyari", o["sorun"]), unsafe_allow_html=True)
        return
    _ad = o["firma"].get("firma_adi") or ""
    st.caption("Kapatınca Ref No Takip'e açılacak: " + " · ".join(
        f"{_ad[:40]} {KR._para(p['tutar'], p['doviz'])} ({p['aciklama']})" for p in o["plan"])
        + ". Kapatan kişiye ve İbrahim'e mail gider.")
    if o["benzer"]:
        st.markdown(mesaj("uyari", "Aynı firma ve ayda elle girilmiş ref var: "
                                   + ", ".join(r.get("ref_no") or "" for r in o["benzer"][:5])
                                   + ". Bu kampanyanın desteği orada da varsa P&L'de iki kez sayılır."),
                    unsafe_allow_html=True)


def _yenile(kid=None, mesaj_metni=None):
    """Kayıttan sonra: önbellek boşalır, (istenirse) pencere yeniden açılır."""
    B.yenile(mesaj_metni, ac=("kmp", kid) if kid else None)


@st.dialog("Kampanya detayı", width="large")
def _detay_dialog(kid):
    from shared.tasarim import yan_panel; yan_panel("orta")   # sağ panel (okuma)
    bugun = tr_today()
    kamp = next((k for k in (get_kampanyalar() or []) if k.get("id") == kid), None)
    if not kamp:
        st.markdown(mesaj("hata", "Kampanya bulunamadı — silinmiş olabilir."), unsafe_allow_html=True)
        return
    urunler_k = get_kampanya_urunler(kid) or []
    _, _, pacal, urunler = _veri()
    o = H.kampanya_ozet(kamp, urunler_k, pacal, bugun)
    ad_d, renk, aciklama = H.DURUMLAR[o["durum"]]
    cipler = "".join(B.cip(t, r) for t, r in (
        (_firma_ad(kamp.get("firma")), "mor2"),
        ((kamp.get("kampanya_turu") or "").strip(), "cyan"),
        (_kat_ad(kamp.get("kategori")), "pembe")) if t)
    st.markdown(_css() + (
        f'<div class="kmp-dt" style="--d:{rv(renk)}">'
        f'<div class="kmp-dt-ust"><span class="kmp-ad">{_h.escape(str(kamp.get("kampanya_adi") or ""))}</span>'
        f'{B.cip(ad_d, renk, aciklama)}{cipler}</div>'
        f'<div class="kmp-zaman"><span>{_aralik(kamp)}</span><div class="kmp-cubuk">'
        f'<i style="width:{(1.0 if o["durum"] == "kapali" else o["oran"]) * 100:.0f}%"></i></div>'
        f'<b>{_kalan_metni(o)}</b></div></div>'), unsafe_allow_html=True)

    kalemler = [
        {"etiket": "Satılan", "deger": f"{tr_sayi(o['adet'])} adet", "renk": "mor",
         "alt": f"{o['urun']} ürün"},
        {"etiket": "Net ciro", "deger": _usd(o["ciro"]), "renk": "cyan", "alt": "destekler düşülmüş",
         "ipucu": f"${tr_sayi(o['ciro'], 2)}"},
        {"etiket": "Destek", "deger": _usd(o["destek"]), "renk": "amber",
         "alt": f"firma {_usd(o['firma_destek'])} · ek {_usd(o['ek_destek'])}"},
    ]
    if o["spiff"] > 0:
        kalemler.append({"etiket": "Spiff", "deger": _usd(o["spiff"]), "renk": "amber",
                         "alt": f"₺{tr_sayi(kamp.get('spiff_tl'))} · {'kesin' if kamp.get('spiff_fatura') else 'tahmini kur'}"})
    kalemler.append({"etiket": "Net kâr", "deger": _usd(o["net"]),
                     "renk": "yesil" if o["net"] >= 0 else "kirmizi",
                     "alt": f"net marj %{tr_sayi(o['marj'], 1)}" if o["marj"] is not None else "adet girilince hesaplanır",
                     "ipucu": f"${tr_sayi(o['net'], 2)}"})
    st.markdown(kpi_serit(kalemler), unsafe_allow_html=True)
    if o["eksik_pacal"]:
        st.markdown(mesaj("uyari", f"{o['eksik_pacal']} üründe paçal maliyet yok; o ürünlerin kârı "
                                   "hesaplanamıyor. Paçal, İthalat kaydı girilince otomatik gelir."),
                    unsafe_allow_html=True)

    t_urun, t_bilgi, t_islem = st.tabs(["Ürünler", "Bilgiler", "İşlemler"])
    with t_urun:
        _sekme_urunler(kamp, urunler_k, pacal, urunler, o)
    with t_bilgi:
        _sekme_bilgiler(kamp)
    with t_islem:
        _sekme_islemler(kamp, o)


def _sekme_urunler(kamp, urunler_k, pacal, urunler, o):
    kid = kamp["id"]
    if urunler_k:
        satirlar = []
        for u in urunler_k:
            h = H.urun_hesap(u, pacal.get(u.get("sku"), 0))
            satirlar.append({
                "_id": u["id"], "SKU": u.get("sku", ""), "Ürün": _urun_ad(u.get("sku"), u.get("urun_adi")),
                "Paçal": h["pacal"] or None, "Satış": h["satis"], "Firma desteği": h["fd"],
                "Ek destek": h["ed"], "Satılan": h["adet"],
                "Net kâr/adet": h["net_kar"], "Net marj %": h["marj"], "Toplam net": h["t_net"] if h["net_kar"] is not None else None,
                "Not": u.get("notlar") or "", "Sil": False,
            })
        df = pd.DataFrame(satirlar).set_index("_id")
        _para = lambda b, y=None: st.column_config.NumberColumn(b, format="dollar", step=0.01, help=y)
        duz = st.data_editor(
            df, key=f"kmp_ed_{kid}", use_container_width=True, hide_index=True,
            height=min(420, 38 + 35 * len(df)),
            disabled=["SKU", "Ürün", "Paçal", "Net kâr/adet", "Net marj %", "Toplam net"],
            column_config={
                "SKU": st.column_config.TextColumn("SKU", width="small"),
                "Ürün": st.column_config.TextColumn("Ürün", width="medium"),
                "Paçal": _para("Paçal", "Kampanyaya eklenirken kaydedilen paçal; yoksa güncel paçal"),
                "Satış": _para("Satış", "Müşteriye birim fiyat"),
                "Firma desteği": _para("Firma desteği", "Firmanın verdiği birim destek"),
                "Ek destek": _para("Ek destek", "Bizim verdiğimiz ek birim destek"),
                "Satılan": st.column_config.NumberColumn("Satılan", min_value=0, step=1, format="localized"),
                "Net kâr/adet": _para("Net kâr/adet", "(Satış − destekler) − paçal"),
                "Net marj %": st.column_config.NumberColumn("Net marj %", format="localized", step=0.1,
                                                            help="Net kâr ÷ net satış"),
                "Toplam net": _para("Toplam net"),
                "Not": st.column_config.TextColumn("Not", width="small"),
                "Sil": st.column_config.CheckboxColumn("Sil", width="small",
                                                       help="İşaretle ve kaydet: ürün kampanyadan çıkar"),
            })
        st.caption("Satış, destek, adet ve notu tabloda düzenle; hesaplı sütunlar kaydedince güncellenir.")
        _degisen, _silinen = [], []
        for rid, yeni in duz.iterrows():
            eski = df.loc[rid]
            if bool(yeni["Sil"]):
                _silinen.append(rid)
            elif any(str(yeni[c]) != str(eski[c]) for c in ("Satış", "Firma desteği", "Ek destek", "Satılan", "Not")):
                _degisen.append((rid, yeni))
        b1, b2, b3 = st.columns([1.3, 1.6, 2])
        _etk = (f"Kaydet ({len(_degisen) + len(_silinen)} değişiklik)" if (_degisen or _silinen)
                else "Kaydet")
        if b1.button(_etk, type="primary", icon=":material/save:", use_container_width=True,
                     disabled=not (_degisen or _silinen), key=f"kmp_kaydet_{kid}"):
            _kaydet_urunler(_degisen, _silinen)
            _yenile(kid, "Kampanya ürünleri kaydedildi")
        if o["durum"] == "bekliyor":
            if b2.button("Kaydet ve kampanyayı kapat", icon=":material/task_alt:", use_container_width=True,
                         key=f"kmp_kapat_hizli_{kid}"):
                _kaydet_urunler(_degisen, _silinen)
                _yenile(kid, _kapat_ve_ref_ac(kamp))
    else:
        st.markdown('<div class="kmp-bos">Bu kampanyada henüz ürün yok. Aşağıdan ekle ya da '
                    'kampanyayı Excel şablonuyla oluştur.</div>', unsafe_allow_html=True)

    # ── Ürün ekle ──
    _var = {u.get("sku") for u in urunler_k}
    _secenek = {f"{u['sku']} — {u.get('urun_adi','')}": u for u in urunler if u["sku"] not in _var}
    with st.expander("Ürün ekle", icon=":material/add_circle:", expanded=not urunler_k):
        if not _secenek:
            st.caption("Eklenebilecek ürün yok (hepsi zaten bu kampanyada ya da ürün kaydı yok).")
            return
        with st.form(f"kmp_urun_ekle_{kid}", clear_on_submit=True, border=False):
            u_sec = st.selectbox("Ürün", list(_secenek.keys()), key=f"kmp_ue_u_{kid}",
                                 index=None, placeholder="SKU ya da ürün adı yaz…")
            e1, e2, e3 = st.columns(3)
            u_satis = e1.number_input("Satış ($)", min_value=0.0, step=0.5, format="%.2f", key=f"kmp_ue_s_{kid}")
            u_fd = e2.number_input("Firma desteği ($)", min_value=0.0, step=0.5, format="%.2f", key=f"kmp_ue_fd_{kid}")
            u_ed = e3.number_input("Ek destek ($)", min_value=0.0, step=0.5, format="%.2f", key=f"kmp_ue_ed_{kid}")
            if st.form_submit_button("Kampanyaya ekle", type="primary", icon=":material/add:"):
                if not u_sec or u_satis <= 0:
                    st.error("Ürün ve satış fiyatı gerekli.")
                else:
                    u = _secenek[u_sec]
                    ekle_kampanya_urun(kid, u["sku"], u.get("urun_adi", u["sku"]),
                                       float(u.get("final_cost_price") or 0), u_satis, u_fd, u_ed, "")
                    _p = float(u.get("final_cost_price") or 0)
                    _ek = ""
                    if _p > 0:
                        _nk = (u_satis - u_fd - u_ed) - _p
                        _ek = f" · birim net kâr ${tr_sayi(_nk, 2)}"
                    _yenile(kid, f"{u['sku']} eklendi{_ek}")


def _kaydet_urunler(degisen, silinen):
    for rid, y in degisen:
        guncelle_kampanya_urun(int(rid), float(y["Satış"] or 0), float(y["Firma desteği"] or 0),
                               float(y["Ek destek"] or 0), int(y["Satılan"] or 0), str(y["Not"] or ""))
    for rid in silinen:
        sil_kampanya_urun(int(rid))


def _sekme_bilgiler(kamp):
    kid = kamp["id"]
    _turler = ["(Seçilmedi)"] + H.KAMPANYA_TURLERI
    _tur = (kamp.get("kampanya_turu") or "").strip()
    if _tur and _tur not in _turler:
        _turler.append(_tur)
    _fs = _firma_secenekleri([kamp])
    from shared.utils import firma_kanonik as _fkn
    _firma = next((f for f in _fs if _fkn(f) == _fkn(kamp.get("firma"))), _fs[-1])   # eski KANAL → DİĞER
    with st.form(f"kmp_bilgi_{kid}", border=False):
        a1, a2 = st.columns([2, 1])
        ad = a1.text_input("Kampanya adı", value=kamp.get("kampanya_adi") or "")
        firma = a2.selectbox("Firma", _fs, index=_fs.index(_firma), format_func=_firma_ad)
        b1, b2, b3, b4 = st.columns(4)
        bas = b1.date_input("Başlangıç", value=H._tarih(kamp.get("baslangic_tarihi")) or tr_today(), format="DD.MM.YYYY")
        bit = b2.date_input("Bitiş", value=H._tarih(kamp.get("bitis_tarihi")) or tr_today(), format="DD.MM.YYYY")
        tur = b3.selectbox("Tür", _turler, index=_turler.index(_tur) if _tur in _turler else 0)
        kat = b4.text_input("Kategori", value=kamp.get("kategori") or "", placeholder="Genel / karışık")
        st.markdown('<div class="kmp-bos" style="margin:6px 0 2px">Spiff: TL tutar ve kur. Fatura gelince '
                    'kuru güncelle, USD maliyet ve kâr kendiliğinden düzelir.</div>', unsafe_allow_html=True)
        s1, s2, s3 = st.columns([1.2, 1, 1.2], vertical_alignment="bottom")
        sp_tl = s1.number_input("Spiff (₺)", min_value=0.0, step=100.0, format="%.2f",
                                value=float(kamp.get("spiff_tl") or 0))
        sp_kur = s2.number_input("Kur (₺/$)", min_value=0.0, step=0.1, format="%.4f",
                                 value=float(kamp.get("spiff_kur") or 0))
        sp_fat = s3.checkbox("Fatura geldi (kur kesin)", value=bool(kamp.get("spiff_fatura")))
        notlar = st.text_area("Notlar", value=kamp.get("notlar") or "", height=80)
        if st.form_submit_button("Bilgileri kaydet", type="primary", icon=":material/save:"):
            if not ad.strip():
                st.error("Kampanya adı boş olamaz.")
            elif bit < bas:
                st.error("Bitiş tarihi başlangıçtan önce olamaz.")
            else:
                guncelle_kampanya(kid, ad.strip(), firma, str(bas), str(bit), notlar,
                                  kategori=kat.strip(), kampanya_turu="" if tur.startswith("(") else tur,
                                  spiff_tl=sp_tl, spiff_kur=sp_kur, spiff_fatura=sp_fat)
                _yenile(kid, "Kampanya bilgileri kaydedildi")


def _sekme_islemler(kamp, o):
    kid = kamp["id"]
    # Kapat / yeniden aç
    if o["durum"] != "kapali":
        st.markdown("**Kampanyayı kapat**")
        st.caption("Kapanan kampanya Kapalı görünümüne geçer; net kâr satılan adetlerle kesinleşir.")
        _adetsiz = o["adet_yok"] and o["urun"] > 0
        if _adetsiz:
            st.markdown(mesaj("uyari", "Hiçbir ürüne satış adedi girilmemiş. Önce Ürünler sekmesinde "
                                       "adetleri gir; adetsiz kapatırsan kâr $0 görünür."), unsafe_allow_html=True)
        onay = st.checkbox("Adet girmeden kapat", key=f"kmp_kapat_onay_{kid}") if _adetsiz else True
        _ref_onizleme(kamp)
        if st.button("Kampanyayı kapat", icon=":material/task_alt:", disabled=not onay, key=f"kmp_kapat_{kid}"):
            _yenile(kid, _kapat_ve_ref_ac(kamp))
    else:
        st.markdown("**Kampanyayı yeniden aç**")
        st.caption("Adet ya da fiyat düzeltmek için kampanyayı yeniden açabilirsin.")
        if st.button("Yeniden aç", icon=":material/lock_open:", key=f"kmp_yac_{kid}"):
            _yeniden_ac(kid)
            _yenile(kid, "Kampanya yeniden açıldı")

    st.divider()
    st.markdown("**Kopyala**")
    st.caption("Ürünleri ve birim destekleriyle yeni bir kampanya oluşturur; satılan adetler sıfırlanır.")
    k1, k2 = st.columns([2.2, 1], vertical_alignment="bottom")
    yeni_ad = k1.text_input("Yeni kampanya adı", value=f"{kamp.get('kampanya_adi','')} (kopya)",
                            key=f"kmp_kop_ad_{kid}")
    if k2.button("Kopyala", icon=":material/content_copy:", use_container_width=True, key=f"kmp_kop_{kid}"):
        yid = ekle_kampanya(yeni_ad.strip() or f"{kamp.get('kampanya_adi','')} (kopya)",
                            kamp.get("firma", ""), kamp.get("baslangic_tarihi"), kamp.get("bitis_tarihi"),
                            kamp.get("notlar", "") or "", kamp.get("kategori", "") or "",
                            kampanya_turu=kamp.get("kampanya_turu", "") or "",
                            spiff_tl=kamp.get("spiff_tl") or 0, spiff_kur=kamp.get("spiff_kur") or 0,
                            spiff_fatura=bool(kamp.get("spiff_fatura")))
        if yid:
            for u in get_kampanya_urunler(kid) or []:
                ekle_kampanya_urun(yid, u.get("sku", ""), u.get("urun_adi", ""), u.get("pacal_maliyet", 0),
                                   u.get("satis_fiyati", 0), u.get("birim_firma_destek", 0),
                                   u.get("birim_ek_destek", 0), u.get("notlar", "") or "")
            _yenile(yid, "Kampanya kopyalandı — yeni kampanya açıldı")
        else:
            st.error("Kopya oluşturulamadı (veritabanı kayıt döndürmedi).")

    st.divider()
    st.markdown("**Sil**")
    if B.onayli_sil(f"Evet, '{kamp.get('kampanya_adi','')}' kampanyasını sil", key=f"kmp_{kid}",
                    dugme="Kampanyayı sil",
                    aciklama="Kampanya ve tüm ürün satırları kalıcı olarak silinir; geri alınamaz."):
        sil_kampanya(kid)
        st.session_state.pop("_kmp_sec", None)
        _yenile(None, "Kampanya silindi")


# ════════════════════════════════════════════════════════════════════
# Yeni kampanya
# ════════════════════════════════════════════════════════════════════
@st.dialog("Yeni kampanya", width="large")
def _yeni_dialog(katlar):
    st.caption("Temel bilgileri gir; kampanya açılınca ürünleri doğrudan ekleyebilirsin.")
    with st.form("kmp_yeni_form", border=False):
        ad = st.text_input("Kampanya adı", placeholder="örn. Hepsiburada Kasım Monitör Haftası")
        a1, a2, a3 = st.columns(3)
        firma = a1.selectbox("Firma", _firma_secenekleri(), index=None, placeholder="Seç…", format_func=_firma_ad)
        tur = a2.selectbox("Tür", H.KAMPANYA_TURLERI, index=None, placeholder="Seç (isteğe bağlı)")
        kat = a3.selectbox("Kategori", katlar, index=None, placeholder="Genel / karışık",
                           format_func=_kat_ad)
        b1, b2 = st.columns(2)
        bas = b1.date_input("Başlangıç", value=tr_today(), format="DD.MM.YYYY")
        bit = b2.date_input("Bitiş", value=tr_today(), format="DD.MM.YYYY")
        if st.form_submit_button("Oluştur ve ürün ekle", type="primary", icon=":material/arrow_forward:"):
            if not ad.strip():
                st.error("Kampanya adı gerekli.")
            elif not firma:
                st.error("Firma seç.")
            elif bit < bas:
                st.error("Bitiş tarihi başlangıçtan önce olamaz.")
            else:
                try:
                    yid = ekle_kampanya(ad.strip(), firma, str(bas), str(bit), "", kat or "",
                                        kampanya_turu=tur or "", spiff_tl=0, spiff_kur=0, spiff_fatura=False)
                except Exception as e:  # noqa: BLE001 — kullanıcıya göster, kaydet
                    from shared.hata_log import kaydet
                    kaydet("kampanya.ekle", e)
                    yid = None
                if yid:
                    _yenile(yid, f"'{ad.strip()}' oluşturuldu")
                else:
                    st.error("Kampanya kaydedilemedi; ayrıntı Sistem Kayıtları'nda.")


# ════════════════════════════════════════════════════════════════════
# Excel şablonundan kampanya — Dosya kapısı (Ekim 2026; eskiden bu ekrandaki "Excel'den" penceresi)
# ════════════════════════════════════════════════════════════════════
_KMP_TAM_KOL = ["FİRMA ADI", "MARKA", "KATEGORİ", "STOK KODU", "STOK ADI", "BARKOD",
                "FİYAT", "REBATE", "SELLOUT", "EK SELLOUT", "SPIFF", "NET FİYAT",
                "KAMPANYA ADI", "KAMPANYA TÜRÜ", "BAŞLANGIÇ TARİHİ", "BİTİŞ TARİHİ"]


def _knrm(_s):
    _s = str(_s).strip()
    for _a, _b in (("İ", "i"), ("I", "i"), ("ı", "i"), ("Ş", "s"), ("ş", "s"),
                   ("Ğ", "g"), ("ğ", "g"), ("Ü", "u"), ("ü", "u"),
                   ("Ö", "o"), ("ö", "o"), ("Ç", "c"), ("ç", "c")):
        _s = _s.replace(_a, _b)
    return _s.lower()


def sablon_kampanya():
    """Kampanya şablonu: Firma · Kategori · Marka · Kampanya Türü açılır listeli (gizli LISTELER sayfası)."""
    kamps, _ku, _pc, urun_data_k = _veri()
    _katlar = _kategoriler(urun_data_k, kamps)
    _tbuf = BytesIO()
    try:
        from satis.database import get_kanallar as _get_cariler
        _cariler = [c for c in (_get_cariler() or []) if str(c).strip()]
    except Exception:
        _cariler = []
    from shared.ana_veri import marka_ad as _marka_ad
    _markalar = sorted({_marka_ad(u.get("marka")) for u in urun_data_k} - {""}, key=str.upper)   # tekil, tek yazım
    _turler = [t for t in H.KAMPANYA_TURLERI if not str(t).startswith("(")]
    try:
        import openpyxl
        from openpyxl.worksheet.datavalidation import DataValidation
        _wb = openpyxl.Workbook()
        _ws = _wb.active
        _ws.title = "Kampanya"
        _ws.append(_KMP_TAM_KOL)
        _lst = _wb.create_sheet("LISTELER")

        def _lst_yaz(_ci, _bas, _veri):
            _lst.cell(row=1, column=_ci, value=_bas)
            for _i, _v in enumerate(_veri, start=2):
                _lst.cell(row=_i, column=_ci, value=str(_v))
            return len(_veri)
        _nf = _lst_yaz(1, "FIRMA", _cariler)
        _nk = _lst_yaz(2, "KATEGORI", _katlar)
        _nm = _lst_yaz(3, "MARKA", _markalar)
        _nt = _lst_yaz(4, "TUR", _turler)
        _lst.sheet_state = "hidden"

        def _ekle_dv(_kol, _lcol, _n):
            if _n <= 0:
                return
            _dv = DataValidation(type="list", formula1=f"=LISTELER!${_lcol}$2:${_lcol}${_n + 1}",
                                 allow_blank=True)
            _dv.add(f"{_kol}2:{_kol}2000")
            _ws.add_data_validation(_dv)
        _ekle_dv("A", "A", _nf)   # FİRMA ADI
        _ekle_dv("B", "C", _nm)   # MARKA (yeni düzende 2. kolon)
        _ekle_dv("C", "B", _nk)   # KATEGORİ (yeni düzende 3. kolon)
        _ekle_dv("N", "D", _nt)   # KAMPANYA TÜRÜ (SPIFF eklendi → N kolonu)
        _wb.save(_tbuf)
    except Exception:
        # Açılır liste kurulamazsa düz şablona düş
        _tbuf = BytesIO()
        with pd.ExcelWriter(_tbuf, engine="openpyxl") as _w:
            pd.DataFrame(columns=_KMP_TAM_KOL).to_excel(_w, index=False, sheet_name="Kampanya")
    return _tbuf.getvalue(), "KAMPANYA_OLUSTUR_SABLONU.xlsx"


def urunleri_ekle(kampanya_id, satirlar, ekle=None):
    """Şablondaki ürünleri kampanyaya ekler. Döner: (eklenen, [{SKU, Hata}]).
    Eskiden ürün eklenemediğinde hata sessizce yutuluyor, kampanya eksik ürünle oluşuyordu."""
    ekle = ekle or ekle_kampanya_urun
    n, hatalar = 0, []
    for _u in satirlar:
        try:
            ekle(kampanya_id, _u["sku"], _u["urun_adi"], _u["pacal"], _u["satis"], _u["fd"], _u["ed"], "")
            n += 1
        except Exception as e:  # noqa: BLE001
            hatalar.append({"SKU": _u["sku"], "Hata": f"{type(e).__name__}: {str(e)[:120]}"})
    return n, hatalar


def kapi_kampanya(dosya, kapi):
    """Tek şablonla YENİ kampanya oluşturur + ürünleri ekler."""
    kamps, _ku, _pc, urun_data_k = _veri()
    _kt_kat_list = _kategoriler(urun_data_k, kamps)
    FIRMA_LISTESI_K = _firma_secenekleri()
    KAMPANYA_TURLERI = ["(Seçilmedi)"] + H.KAMPANYA_TURLERI
    urun_dict_k = {u["sku"]: u for u in urun_data_k}
    st.caption("Firma · Kampanya Türü · Başlangıç/Bitiş Tarihi · Ek Sellout dahil tüm bilgiler dosyadan gelir. "
               "Eşleme: FİYAT→satış · SELLOUT→firma desteği · EK SELLOUT→ek destek (paçal SKU'dan otomatik).")
    try:
        _kdf = pd.read_excel(dosya)
    except Exception as _e:
        st.error(f"Excel okunamadı: {_e}")
        return
    _kolmap = {_knrm(c): c for c in _kdf.columns}

    def _col(*_names):
        _ns = [_knrm(n) for n in _names]
        for _n in _ns:  # önce tam eşleşme (fiyat vs net fiyat, sellout vs ek sellout)
            if _n in _kolmap:
                return _kolmap[_n]
        for _n in _ns:  # sonra alt-dizge
            for _k, _actual in _kolmap.items():
                if _n in _k:
                    return _actual
        return None

    def _gv(_row, *_names):
        _c = _col(*_names)
        if _c is not None and _c in _row and pd.notna(_row[_c]):
            return _row[_c]
        return None

    def _knum(_v):
        try:
            if _v is None or (isinstance(_v, float) and pd.isna(_v)):
                return 0.0
            return float(str(_v).replace(",", ".").replace(" ", ""))
        except Exception:
            return 0.0

    def _tarih(_v):
        from shared.utils import tarih_metni
        return tarih_metni(_v)

    st.dataframe(_kdf.head(15), use_container_width=True, height=200)
    _xl_ad = _xl_firma = _xl_kat = _xl_tur = ""
    _xl_bas = _xl_bit = None
    _xl_spiff = 0.0
    for _, _r in _kdf.iterrows():
        _xl_ad = _xl_ad or str(_gv(_r, "kampanya adi") or "").strip()
        _xl_firma = _xl_firma or str(_gv(_r, "firma adi", "firma") or "").strip()
        _xl_kat = _xl_kat or str(_gv(_r, "kategori") or "").strip()
        _xl_tur = _xl_tur or str(_gv(_r, "kampanya turu", "tur") or "").strip()
        _xl_bas = _xl_bas or _tarih(_gv(_r, "baslangic"))
        _xl_bit = _xl_bit or _tarih(_gv(_r, "bitis"))
        if not _xl_spiff:
            _xl_spiff = _knum(_gv(_r, "spiff"))
        if _xl_ad and _xl_firma and _xl_tur and _xl_bas and _xl_bit and _xl_spiff:
            break
    # Kategori: mevcut bir kategoriyle yalnızca harf farkı varsa onu kullan (KASA/kasa mükerrerini önle)
    _kat_final = _xl_kat
    for _ek in _kt_kat_list:
        if _ek.strip().lower() == _xl_kat.lower():
            _kat_final = _ek
            break
    _tur_idx = 0
    for _ti, _t in enumerate(KAMPANYA_TURLERI):
        if _t.lower() == _xl_tur.lower():
            _tur_idx = _ti
            break

    # Firma: şablondaki TAM cari adını FIRMA_LISTESI_K koduna (HB/VATAN/ITOPYA...) eşle.
    # Kampanya 'firma' alanı KOD saklar; ekranda firma_gorunen_ad ile tam ad gösterilir.
    def _firma_koda(_ad):
        _adn = _knrm(_ad)
        if not _adn:
            return ""
        for _c in FIRMA_LISTESI_K:
            try:
                _tam = _knrm(firma_gorunen_ad(_c, kisa=False))  # eşleştirme: TAM ad
            except Exception:
                _tam = ""
            if _tam and (_tam == _adn or _tam in _adn or _adn in _tam):
                return _c
        for _anahtar, _c in (("d-market", "HB"), ("dmarket", "HB"), ("hepsiburada", "HB"),
                             ("eera", "ITOPYA"), ("itopya", "ITOPYA"),
                             ("vatan", "VATAN"), ("monday", "MONDAY")):
            if _anahtar in _adn:
                return _c
        return ""
    _fk = _firma_koda(_xl_firma)
    _firma_idx = FIRMA_LISTESI_K.index(_fk) if _fk in FIRMA_LISTESI_K else 0
    # Anahtarlar dosyaya özgü: ikinci dosyada ilk dosyanın değerleri kalmasın (eskiden kalıyordu)
    _of1, _of2 = st.columns(2)
    _o_ad = _of1.text_input("Kampanya Adı *", value=_xl_ad, key=kapi.anahtar("kmp_o_ad"))
    _o_firma = _of2.selectbox("Firma *", FIRMA_LISTESI_K, index=_firma_idx,
                              format_func=firma_gorunen_ad, key=kapi.anahtar("kmp_o_firma"))
    if _xl_firma:
        st.caption(f"Dosyadaki firma: **{_xl_firma[:44]}** → "
                   + (f"**{_o_firma}** koduyla eşlendi." if _fk
                      else "otomatik eşlenemedi — yukarıdan doğru firmayı seç."))
    _of3, _of4, _of5 = st.columns(3)
    _o_turu = _of3.selectbox("Kampanya Türü", KAMPANYA_TURLERI, index=_tur_idx, key=kapi.anahtar("kmp_o_turu"))
    _o_spiff_tl, _o_spiff_kur = 0.0, 0.0
    if str(_o_turu).lower() == "spiff" or _xl_spiff > 0:
        _sp1, _sp2 = st.columns(2)
        _o_spiff_tl = _sp1.number_input("Spiff Net (₺TL)", min_value=0.0, step=100.0,
                                        value=float(_xl_spiff or 0), format="%.4f",
                                        key=kapi.anahtar("kmp_o_spiff_tl"),
                                        help="Dosyadaki SPIFF kolonundan geldi; düzenleyebilirsin.")
        _o_spiff_kur = _sp2.number_input("Kur (₺/$ tahmini)", min_value=0.0, step=0.1,
                                         value=None, placeholder="örn. 40",
                                         format="%.4f", key=kapi.anahtar("kmp_o_spiff_kur")) or 0.0
        if _o_spiff_tl and _o_spiff_kur:
            st.caption(f"≈ ${tr_sayi((_o_spiff_tl / _o_spiff_kur), 2)} USD spiff maliyeti (tahmini)")
    _o_bas = _of4.date_input("Başlangıç Tarihi *", value=(_xl_bas or tr_today()), key=kapi.anahtar("kmp_o_bas"),
                             format="DD.MM.YYYY")
    _o_bit = _of5.date_input("Bitiş Tarihi *", value=(_xl_bit or tr_today()), key=kapi.anahtar("kmp_o_bit"),
                             format="DD.MM.YYYY")
    st.caption(f"Kategori (dosyadan): **{_kat_final or '—'}**"
               + ("" if (not _kat_final or _kat_final == _xl_kat)
                  else f"  · mevcut '{_kat_final}' ile eşleştirildi (yeni mükerrer açılmadı)"))
    _urun_satir = []
    for _, _r in _kdf.iterrows():
        _sku = str(_gv(_r, "stok kodu", "sku") or "").strip()
        if not _sku or _sku.lower() == "nan":
            continue
        _bilgi = urun_dict_k.get(_sku, {})
        _uad = (str(_gv(_r, "stok adi", "urun adi", "urun") or "").strip() or _bilgi.get("urun_adi", _sku))
        _urun_satir.append({"sku": _sku, "urun_adi": _uad, "pacal": _bilgi.get("final_cost_price", 0) or 0,
                            "satis": _knum(_gv(_r, "fiyat", "satis")), "fd": _knum(_gv(_r, "sellout")),
                            "ed": _knum(_gv(_r, "ek sellout"))})
    st.caption(f"Şablonda **{len(_urun_satir)}** geçerli ürün satırı bulundu.")
    # Aynı şablon ikinci kez yüklenirse ikinci bir kampanya açılırdı (destekler iki kez sayılır)
    _ayni = [k for k in kamps if normalize_tr(str(k.get("kampanya_adi") or "")) == normalize_tr(_o_ad)
             and normalize_tr(str(k.get("firma") or "")) == normalize_tr(_o_firma)
             and str(k.get("baslangic_tarihi") or "")[:10] == str(_o_bas)]
    _ayni_onay = True
    if _ayni and _o_ad.strip():
        st.warning(f"**{_o_ad.strip()}** adında, aynı firma ve başlangıç tarihli bir kampanya zaten var. "
                   "Aynı şablonu ikinci kez yüklersen ikinci bir kampanya açılır.")
        _ayni_onay = st.checkbox("Bu farklı bir kampanya — yine de oluştur", key=kapi.anahtar("kmp_o_ayni"))
    if st.button("Şablondan Kampanya Oluştur ve Ürünleri Ekle", type="primary", use_container_width=True,
                 key=kapi.anahtar("kmp_o_olustur"),
                 disabled=(not _o_ad.strip() or not _o_firma.strip() or not _urun_satir or not _ayni_onay),
                 icon=":material/rocket_launch:"):
        _ohata, _oyid = None, None
        try:
            _oyid = ekle_kampanya(
                _o_ad.strip(), _o_firma.strip(), str(_o_bas), str(_o_bit), "", _kat_final,
                kampanya_turu=("" if str(_o_turu).startswith("(") else _o_turu),
                spiff_tl=(_o_spiff_tl or 0), spiff_kur=(_o_spiff_kur or 0))
        except Exception as _e:
            _ohata = str(_e)
        if not _oyid:
            st.error(f"Kampanya oluşturulamadı — {_ohata}" if _ohata
                     else "Kampanya oluşturulamadı (tablo izni/kolon olabilir).")
            return
        _on, _uhata = urunleri_ekle(_oyid, _urun_satir)
        from shared.yukleme_gecmisi import kaydet as _yg_kaydet
        _yg_kaydet("kampanya_sablon", _on, dosya.name)
        st.cache_data.clear()
        kapi.bitti(f"'{_o_ad.strip()}' kampanyası oluşturuldu ve {_on} ürün eklendi "
                   f"(Firma: {_o_firma.strip()} · Tür: {_o_turu} · {_o_bas}→{_o_bit})."
                   + (f" {len(_uhata)} ürün eklenemedi; liste aşağıda — kampanya detayından elle ekleyebilirsin."
                      if _uhata else ""), tablo=_uhata, uyari=bool(_uhata))
