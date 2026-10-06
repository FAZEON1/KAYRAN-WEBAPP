# -*- coding: utf-8 -*-
"""Sayfa testi · Dosya kapısı senaryoları (Ekim 2026, kapı canlıya çıktıktan sonraki tarama).

Kapının içindeki 18 akış, gerçek Streamlit'te "olmaması gereken" durumlarla denendi:
- Elle YANLIŞ tür seçilirse (dosya o türe benzemiyor) kayıt, uyarı onaylanmadan yapılamaz. Taramada
  306 dosya/tür eşleşmesinin 26'sında yanlış dosya kaydedilebiliyordu (ör. sipariş Excel'i "stok
  değeri raporu" seçilince 5. sütundaki sayılar stok değeri diye yazılıyordu).
- Verisiz (yalnız başlık) dosya kaydedilemez. Boş G5F dosyası BÜTÜN ürünlerin depo stoğunu
  sıfırlıyordu; boş Happy Life dosyası o günün kayıtlarını siliyordu.
- .xls (Mikro'nun eski biçimi) dosyalar da baştan sona aynı akıştan kaydedilir.
- Her sayfadan, bekleyen pencere bayrakları varken bile kapı açılır ("tek pencere" hatası yok;
  Kâr / P&L'deki maliyet düzeltme penceresi bayrağı kapıyı açtırmıyordu).
Yalnız gerçek Streamlit kuruluyken çalışır (CI: "Sayfa testi").
"""
import io
import os

import pytest

pytest.importorskip("streamlit.testing.v1")

from test_dosya_kapisi import (KAYDET, ORNEKLER, _anahtar, _kapida, _sorunlar,  # noqa: E402
                               _uygulama)
from test_sayfalar import SAYFALAR  # noqa: E402

_XLS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "kapi_xls")

# Tarama: yanlış türün akışında kayıt düğmesi ETKİN çıkan (dosya, seçilen tür) çiftleri
YANLIS_TUR = [
    ("aktif_cari", "alinan_destek"), ("aktif_cari", "cek_listesi"), ("aktif_ithalat", "alinan_destek"),
    ("aktif_stok", "alinan_destek"), ("cek_listesi", "aktif_ithalat"), ("cek_listesi", "alinan_destek"),
    ("g5f_sayim", "alinan_destek"), ("gider_tablosu", "aktif_stok"), ("gider_tablosu", "alinan_destek"),
    ("gider_tablosu", "odeme_listesi"), ("happylife", "alinan_destek"), ("iade_excel", "aktif_stok"),
    ("iade_excel", "alinan_destek"), ("ithalat_rapor", "alinan_destek"), ("kampanya_sablon", "alinan_destek"),
    ("mikro_fatura", "alinan_destek"), ("musteri_haftalik", "alinan_destek"), ("odeme_listesi", "alinan_destek"),
    ("ref_excel", "alinan_destek"), ("ref_excel", "cek_listesi"), ("siparis_itopya", "aktif_stok"),
    ("siparis_itopya", "alinan_destek"), ("siparis_vatan", "aktif_stok"), ("siparis_vatan", "alinan_destek"),
    ("siparis_vatan", "g5f_sayim"), ("toplu_mal_kabul", "alinan_destek"),
]


@pytest.mark.parametrize("dosya_tur, secilen", YANLIS_TUR, ids=[f"{a}-{b}" for a, b in YANLIS_TUR])
def test_yanlis_tur_onaysiz_kaydedilmez(dosya_tur, secilen):
    from shared.dosya_tani import tani
    ad, veri = ORNEKLER[dosya_tur]()
    at = _kapida(secilen, ad=ad, veri=veri, adaylar=[a for a in tani(ad, veri) if a.get("tur")])
    assert not _sorunlar(at), _sorunlar(at)
    assert not [b for b in at.button if b.key == _anahtar(KAYDET[secilen])], "uyarı onaylanmadan kayıt düğmesi var"
    assert any("benzemiyor" in w.value for w in at.warning)
    at.checkbox(key="kapi_tur_onay_t1").check().run()            # bilerek devam: akış görünür
    # Akış dosyayı kendi kurallarıyla okur: hata KUTUSU beklenebilir (yanlış dosya), çökme olmamalı
    cokme = [e.value for e in at.exception] + list(at.session_state["_duman_hatalar"]
                                                   if "_duman_hatalar" in at.session_state else [])
    assert not cokme, cokme


def _yalniz_baslik(veri):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(veri))
    for ws in wb.worksheets:
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


BOS = ["g5f_sayim", "happylife", "ref_excel", "alinan_destek", "ithalat_rapor", "musteri_haftalik",
       "siparis_vatan", "siparis_itopya", "mikro_fatura", "iade_excel", "kampanya_sablon", "toplu_mal_kabul"]


@pytest.mark.parametrize("tur", BOS)
def test_verisiz_dosya_kaydedilmez(tur):
    """Yalnız başlık satırı olan dosya: kayıt düğmesi yok ya da pasif (boş G5F stoğu sıfırlıyordu)."""
    ad, veri = ORNEKLER[tur]()
    at = _kapida(tur, ad=ad, veri=_yalniz_baslik(veri))
    if tur == "siparis_itopya":
        at.selectbox(key=_anahtar("sg_kanal")).select_index(0).run()
        at.text_input(key=_anahtar("sg_sno")).input("S-1").run()
    if tur == "ref_excel":
        at.selectbox(key=_anahtar("ref_firma")).select_index(0).run()
    assert not _sorunlar(at), _sorunlar(at)
    d = [b for b in at.button if b.key == _anahtar(KAYDET[tur])]
    assert not d or d[0].disabled, "verisiz dosyada kayıt düğmesi etkin"


@pytest.mark.parametrize("tur", sorted(ORNEKLER))
def test_xls_dosyasi_bastan_sona(tur):
    """Mikro'nun .xls (Excel 97) çıktıları: tanınır, akış çizilir, kayıt sonucu kapıya döner."""
    pytest.importorskip("xlrd")
    from shared.dosya_tani import tani
    from test_dosya_kapisi import test_kapida_akis_ve_kayit
    veri = open(os.path.join(_XLS, f"{tur}.xls"), "rb").read()
    assert [a["tur"] for a in tani(f"{tur}.xls", veri)] == [tur]
    test_kapida_akis_ve_kayit(tur, ad=f"{tur}.xls", veri=veri)


BAYRAKLAR = ["_mk_dialog_ac", "_ms_dialog_ac", "_mlyt_ac", "_talep_ac", "_iade_ac", "_ref_ac", "_kmp_ac",
             "_satis_ac", "_urun_ac", "_sip_ac", "_ith_ac", "_ts_ac", "_odeme_ac"]


def _kimlik(p):
    mod, _, s = p
    return f"{mod}/{s[1]}" if s else mod


@pytest.mark.parametrize("sayfa", SAYFALAR, ids=[_kimlik(p) for p in SAYFALAR])
def test_her_sayfadan_bekleyen_pencere_bayraklariyla_acilir(sayfa):
    mod, anahtar, s = sayfa
    at = _uygulama(mod, **({anahtar: s[0]} if anahtar else {}))
    assert not _sorunlar(at), _sorunlar(at)
    for b in BAYRAKLAR:
        at.session_state[b] = True
    [b for b in at.button if b.key == "ust_dosya"][0].click().run()
    assert not _sorunlar(at), _sorunlar(at)
    assert at.session_state["_kapi_acik"]
    assert not [b for b in BAYRAKLAR if b in at.session_state and at.session_state[b]]
    # Kapı gerçekten açık: sayfa st.stop() ile dursa da (Ödemeler "veri yok", P&L) pencere çizilir
    assert any("Programın kullandığı bütün Excel" in c.value for c in at.caption), "kapı penceresi çizilmedi"


@pytest.mark.parametrize("ekran", ["liste", "govde"])
def test_her_ekranda_kapat_dugmesi(ekran):
    """Pencere X ile kapanmaz (kaydı yarıda kesmesin); her ekranda Kapat düğmesi kapıyı kapatır."""
    at = _kapida("ref_excel")
    if ekran == "liste":
        at.button(key="kapi_geri").click().run()
    at.button(key="kapi_kapat_ust").click().run()
    assert not _sorunlar(at), _sorunlar(at)
    assert "_kapi_acik" not in at.session_state and "_kapi_dosyalar" not in at.session_state


def test_gider_yili_bulunamazsa_secilmeden_kaydedilmez():
    import pandas as pd
    from kapi_ornekleri import AYLAR, _xlsx
    satirlar = [["GİDER TABLOSU"] + [None] * 13, ["KATEGORİ", "KALEM"] + AYLAR,
                ["SABİT GİDERLER", None] + [None] * 12, ["Sabit", "Kira"] + [100000] * 12]
    at = _kapida("gider_tablosu", ad="gider.xlsx", veri=_xlsx({"Gider": pd.DataFrame(satirlar)}, baslik=False))
    assert not _sorunlar(at), _sorunlar(at)
    assert not [b for b in at.button if b.key == _anahtar("gider_kaydet")]
    assert any("yıl bulunamadı" in i.value for i in at.info)
    at.selectbox(key=_anahtar("gider_yil")).select_index(1).run()
    assert [b for b in at.button if b.key == _anahtar("gider_kaydet")]


def test_ayni_kampanya_ikinci_kez_onaysiz_olusmaz():
    k = {"id": 1, "kampanya_adi": "Ekim fırsat", "firma": "VATAN", "baslangic_tarihi": "2026-10-01",
         "bitis_tarihi": "2026-10-31", "durum": "aktif", "olusturma_tarihi": "2026-10-01"}
    at = _kapida("kampanya_sablon", tablolar={"kampanyalar": [k]})
    assert not _sorunlar(at), _sorunlar(at)
    assert at.button(key=_anahtar("kmp_o_olustur")).disabled
    at.checkbox(key=_anahtar("kmp_o_ayni")).check().run()
    assert not at.button(key=_anahtar("kmp_o_olustur")).disabled


def test_fatura_bazli_iade_dokumu_kapidan_yuklenir():
    """Mikro fatura bazlı iade dökümü kapıda iade olarak açılır, kaydedilir; depo dosyadan gelir."""
    from kapi_ornekleri import iade_fatura
    from shared.dosya_tani import tani
    import sahte_db
    ad, veri = iade_fatura()
    at = _kapida("iade_excel", ad=ad, veri=veri, adaylar=[a for a in tani(ad, veri) if a.get("tur")])
    assert not _sorunlar(at), _sorunlar(at)
    assert any("Fatura bazlı iade dökümü" in c.value for c in at.caption)
    assert any("İADE DEPO 4 adet" in c.value and "MERKEZ DEPO 23 adet" in c.value for c in at.caption)
    at.button(key=_anahtar(KAYDET["iade_excel"])).click().run()
    assert not _sorunlar(at), _sorunlar(at)
    # Aynı SKU + cari + depo toplanır: X24F165S iki faturada (2 + 1) → tek satır; SKU kart yazımıyla (FAZE2)
    assert at.session_state["_kapi_sonuc"]["mesaj"].startswith("3 iade kaydedildi")
    yaz = [r for r in sahte_db.TABLOLAR["iadeler"] if r.get("kaynak") == "excel" and r.get("tarih") == "2026-09-30"]
    assert sorted((r["sku"], r["iade_adet"], r["depo"], r["donem_bas"]) for r in yaz) == sorted([
        ("FAZE2", 23, "MERKEZ DEPO", "2026-07-24"), ("X24F165S", 3, "İADE DEPO", "2026-07-24"),
        ("VG27AQ", 1, "İADE DEPO", "2026-07-24")])
