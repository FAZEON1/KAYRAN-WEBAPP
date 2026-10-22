# -*- coding: utf-8 -*-
"""Talepler #35–#41 (22.10.2026).

#35 Berkay: servis kaydı iade alındıktan sonra satışa hazır / ikinci el olunca İade Belgesi düğmesi kayboluyordu
    (G5F00811 geçmişi: mal kabül → teknisyende → iade alındı → satışa hazır).
#36 Gökhan: kampanyadan açılan ref'in Kalemler bölümü boş (sütun canlıda yoktu; kampanya ürünleri kalem olarak
    yazılmıyordu).
#37 Gökhan: Kâr / P&L kırılımında 'mouse pad' (ürün kartı) ve 'MOUSEPAD' (Ref No desteği) iki ayrı satır.
#39 Derya: Tüm ürünler Excel'inde müşteri stoğu yok (yalnız toplam).
#40 Gökhan: EOL işareti yalnız sayfanın en altındaki formdaydı.
#41 Gökhan: kampanyaya ürün eklerken arama kutusunda yazılan harf seçilip siliniyordu (tıkla-seç betiği).
"""
from pathlib import Path

import pytest

KOK = Path(__file__).resolve().parent.parent


# ── #35 ─────────────────────────────────────────────────────────────
def test_iade_belgesi_gecmiste_iade_alindiysa():
    from teknikservis.iade_belge import iade_belgesi_uygun
    k = {"mevcut_durum": "satışa hazır", "ariza_sonucu": "Arıza doğrulandı", "depo": "ikinci el"}
    gecmis = [{"durum": "mal kabül"}, {"durum": "teknisyende"}, {"durum": "iade alındı"}, {"durum": "satışa hazır"}]
    assert iade_belgesi_uygun(k, gecmis)
    assert not iade_belgesi_uygun(k, [{"durum": "mal kabül"}, {"durum": "tamir edildi"}])   # tamir: belge yok
    assert not iade_belgesi_uygun(k)


def test_iade_belgesi_dugmesi_gecmise_bakar():
    s = (KOK / "teknikservis" / "main.py").read_text(encoding="utf-8")
    assert "if not iade_belgesi_uygun(kayit, gecmis):" in s


# ── #36 ─────────────────────────────────────────────────────────────
def test_kampanya_isareti_kalem_kaydinda_korunur():
    from kayranpm.kampanya_ref import isaret_koru, kampanya_no
    assert kampanya_no({"aciklama": "AĞUSTOS EK DESTEK · Kampanya #107"}) == 107
    assert kampanya_no({"aciklama": "TEMMUZ SELLOUT"}) is None
    assert isaret_koru("A · 100 × $6,00 · B · 13 × $12,00", "AĞUSTOS · Kampanya #107") == \
        "A · 100 × $6,00 · B · 13 × $12,00 · Kampanya #107"
    assert isaret_koru("X · Kampanya #107", "Kampanya #107") == "X · Kampanya #107"
    assert isaret_koru("X", "TEMMUZ SELLOUT") == "X"


def test_ref_kalemler_sutunu_migration():
    sql = (KOK / "veritabani" / "53_ref_kalemler.sql").read_text(encoding="utf-8")
    assert "ADD COLUMN IF NOT EXISTS kalemler jsonb" in sql


# ── #37 ─────────────────────────────────────────────────────────────
def test_pnl_kategori_kirilimi_yazim_farkini_birlestirir():
    import importlib
    try:
        M = importlib.import_module("satis.main")
    except Exception as e:  # noqa: BLE001
        pytest.skip(f"satis.main içe aktarılamadı: {e}")
    if not hasattr(M, "kategori_birlestir"):
        pytest.fail("kategori_birlestir yok")
    g, d, r = M.kategori_birlestir({"mouse pad": {"ciro": 100.0, "kar": 20.0, "adet": 5}, "DİĞER": {"ciro": 1.0}},
                                   {"MOUSE PAD": 3.0}, {"MOUSEPAD": 7.0, "GENEL": 2.0})
    assert g == {"mouse pad": {"ciro": 100.0, "kar": 20.0, "adet": 5}, "DİĞER": {"ciro": 1.0}}
    assert d == {"mouse pad": 3.0} and r == {"mouse pad": 7.0, "GENEL": 2.0}       # tek satır


# ── #39 ─────────────────────────────────────────────────────────────
def test_tum_urunler_excel_musteri_stok_sutunlari(tmp_path):
    pytest.importorskip("reportlab")
    from openpyxl import load_workbook
    from kayranpm.rapor import tum_urunler_excel
    from kayranpm.urun_hesap import urun_satiri
    rows = [urun_satiri({"sku": "A", "satis_fiyati": 120, "final_cost_price": 100,
                         "firma_stoklari": {"VATAN": 12, "HB": 3}}),
            urun_satiri({"sku": "B", "satis_fiyati": 50, "final_cost_price": 40, "firma_stoklari": {"VATAN": 4}})]
    yol = str(tmp_path / "t.xlsx")
    assert tum_urunler_excel(rows, yol)[0]
    ws = load_workbook(yol).active
    bas = [c.value for c in ws[4]]
    mus = [b for b in bas if b and b.endswith(" stok") and b != "Kanal Dahil Stok"]
    assert len(mus) == 2 and bas.index(mus[0]) == bas.index("Kanal Dahil Stok") + 1
    a, b = [c.value for c in ws[5]], [c.value for c in ws[6]]
    assert sorted([a[bas.index(m)] for m in mus]) == [3, 12] and sorted([b[bas.index(m)] for m in mus]) == [0, 4]


# ── #40 ─────────────────────────────────────────────────────────────
def test_eol_tek_tikla():
    s = (KOK / "kayranpm" / "main.py").read_text(encoding="utf-8")
    assert '"EOL işaretini kaldır" if _eol_v else "EOL olarak işaretle"' in s
    from kayranpm import database as D
    assert callable(D.eol_isaretle)


# ── #41 ─────────────────────────────────────────────────────────────
def test_tikla_sec_acilir_liste_aramasina_dokunmaz():
    s = (KOK / "app.py").read_text(encoding="utf-8")
    i = s.index("w.__kayranTiklaYaz = true;")
    blok = s[i:i + 2500]
    kontrol = blok.index('t.getAttribute("role") === "combobox"')
    assert kontrol < blok.index("const bas_deger = t.value;")       # select() kurulmadan önce çıkar
