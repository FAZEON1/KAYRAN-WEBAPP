# -*- coding: utf-8 -*-
"""Stok yaşı — FIFO (kayranpm/stok_yasi.py, Ekim 2026).

Yaş mal depoya girdiği gün (teslim tarihi) başlar; en eski parti önce satılır, elde kalan stok en
yeni partilerdir. Bizim stok: satılabilir depolar; müşteri stoğu: son rapor ↔ o müşteriye satışlar.
"""
from datetime import date

import kayranpm.stok_yasi as Y
from shared.utils import sku_anahtar, normalize_tr

BUGUN = date(2026, 10, 4)


def test_fifo_en_yeniler_elde_kalir():
    p = [{"tarih": "2026-01-10", "adet": 100}, {"tarih": "2026-09-20", "adet": 50},
         {"tarih": "2026-06-01", "adet": 80}]
    kal, kaps = Y.fifo_kalan(p, 100)
    assert [(k["tarih"], k["kalan"]) for k in kal] == [("2026-06-01", 50), ("2026-09-20", 50)]
    assert kaps == 0


def test_fifo_stok_partilerden_fazla_kapsanmayan():
    kal, kaps = Y.fifo_kalan([{"tarih": "2026-09-01", "adet": 30}], 45)
    assert [k["kalan"] for k in kal] == [30] and kaps == 15
    assert Y.fifo_kalan([], 0) == ([], 0.0) and Y.fifo_kalan(None, 10) == ([], 10.0)
    assert Y.fifo_kalan([{"tarih": "", "adet": 5}, {"tarih": "2026-01-01", "adet": 0}], 3) == ([], 3.0)


def test_yas_ozeti_agirlikli_ortalama_ve_gruplar():
    kal = [{"tarih": "2026-06-01", "kalan": 50}, {"tarih": "2026-09-20", "kalan": 50}]
    o = Y.yas_ozeti(kal, 10, BUGUN)
    assert o["en_eski_gun"] == 125 and o["en_eski_tarih"] == "2026-06-01" and o["en_yeni_gun"] == 14
    assert round(o["ort_gun"], 1) == round((50 * 125 + 50 * 14) / 100, 1)
    assert o["gruplar"]["91–180 gün"] == 50 and o["gruplar"]["0–30 gün"] == 50
    assert o["gruplar"][Y.KAYITSIZ] == 10 and o["stok"] == 110 and o["kapsanan"] == 100
    assert Y.yas_ozeti([], 5, BUGUN)["ort_gun"] is None


def test_grup_sinirlari():
    assert [Y.grup_adi(g) for g in (0, 30, 31, 60, 61, 90, 91, 180, 181, 900)] == \
        ["0–30 gün", "0–30 gün", "31–60 gün", "31–60 gün", "61–90 gün", "61–90 gün",
         "91–180 gün", "91–180 gün", "180+ gün", "180+ gün"]


def test_bizim_partiler_yalniz_teslim_alinan_teslim_tarihiyle():
    dosyalar = [{"id": 1, "durum": "Teslim Alındı", "teslim_tarihi": "2026-05-02", "tarih": "2026-03-01",
                 "dosya_no": "D1", "teslim_deposu": "MERKEZ DEPO"},
                {"id": 2, "durum": "Antrepoda", "teslim_tarihi": None, "tarih": "2026-08-01", "dosya_no": "D2"},
                {"id": 3, "durum": "Teslim Alındı", "teslim_tarihi": "2026-07-03", "dosya_no": "YI-1",
                 "alim_turu": "yurtici"}]
    kalemler = [{"dosya_id": 1, "sku": "Fazeon X24", "adet": 10}, {"dosya_id": 2, "sku": "X24", "adet": 99},
                {"dosya_id": 3, "sku": "FAZE1", "adet": 500}, {"dosya_id": 1, "sku": "X24", "adet": 0}]
    p = Y.bizim_partiler(dosyalar, kalemler, sku_anahtar)
    assert p == {"X24": [{"tarih": "2026-05-02", "adet": 10, "belge": "D1", "tur": "ithalat",
                          "depo": "MERKEZ DEPO"}],
                 "FAZE1": [{"tarih": "2026-07-03", "adet": 500, "belge": "YI-1", "tur": "yurtici", "depo": ""}]}


def test_musteri_partileri_cari_adindan_firma_koduna():
    tam = {"VATAN": "VATAN BILGISAYAR SANAYI VE TICARET ANONIM SIRKETI",
           "ITOPYA": "EERA ELEKTRONİK TİCARET VE BİLİŞİM HİZMETLERİ ANONİM ŞİRKETİ"}
    bul = Y.firma_cozucu(["VATAN", "ITOPYA"], lambda k: tam.get(k, k), normalize_tr)
    assert bul("Eera Elektronik Ticaret ve Bilişim Hizmetleri Anonim Şirketi") == "ITOPYA"
    assert bul("VATAN") == "VATAN" and bul("BAŞKA FİRMA") is None
    sat = [{"kanal": tam["VATAN"], "sku": "Fazeon X24", "adet": 20, "tarih": "2026-09-01", "siparis_no": "S1"},
           {"kanal": tam["VATAN"], "sku": "X24", "adet": -3, "tarih": "2026-09-02"},
           {"kanal": "BAŞKA FİRMA", "sku": "X24", "adet": 5, "tarih": "2026-09-03"}]
    assert Y.musteri_partileri(sat, bul, sku_anahtar) == \
        {"VATAN": {"X24": [{"tarih": "2026-09-01", "adet": 20, "belge": "S1"}]}}


def test_urun_yasi_ve_toplam():
    o1, kal = Y.urun_yasi(60, [{"tarih": "2026-09-04", "adet": 40}, {"tarih": "2026-03-08", "adet": 40}], BUGUN)
    assert [k["kalan"] for k in kal] == [20, 40] and o1["en_eski_gun"] == 210
    o2, _ = Y.urun_yasi(5, [], BUGUN)
    t = Y.toplam_ozet([o1, o2])
    assert t["stok"] == 65 and t["kapsanmayan"] == 5 and t["gruplar"]["180+ gün"] == 20
    assert round(t["ort_gun"], 2) == round(o1["ort_gun"], 2)


# ── Sayfa: yaş grubuna tıklayınca ürünler, Excel (Ekim 2026) ────────
def _bizim():
    o1, k1 = Y.urun_yasi(60, [{"tarih": "2026-09-04", "adet": 40, "belge": "D2", "tur": "ithalat"},
                              {"tarih": "2026-03-08", "adet": 40, "belge": "D1", "tur": "ithalat"}], BUGUN)
    o2, k2 = Y.urun_yasi(15, [{"tarih": "2026-09-20", "adet": 10, "belge": "YI-1", "tur": "yurtici"}], BUGUN)
    return {"A": (o1, k1, {"sku": "A", "urun_adi": "Ürün A", "kategori": "Monitör"}),
            "B": (o2, k2, {"sku": "B", "urun_adi": "Ürün B", "kategori": "Mouse pad"})}


def test_grup_secilince_yalniz_o_gruptaki_urunler():
    b, pacal = _bizim(), {"A": 10.0, "B": 2.0}
    r = Y.urun_satirlari(b, pacal, grup="180+ gün")
    assert [(x["SKU"], x["Bu yaştaki adet"], x["Bu yaştaki değer ($)"]) for x in r] == [("A", 20, 200.0)]
    r = Y.urun_satirlari(b, pacal, grup="0–30 gün")
    assert [(x["SKU"], x["Bu yaştaki adet"]) for x in r] == [("A", 40), ("B", 10)]
    r = Y.urun_satirlari(b, pacal, grup=Y.KAYITSIZ)
    assert [(x["SKU"], x["Bu yaştaki adet"]) for x in r] == [("B", 5)]
    tum = Y.urun_satirlari(b, pacal)
    assert len(tum) == 2 and tum[0]["180+ gün"] == 20 and tum[1][Y.KAYITSIZ] == 5


def test_parti_satirlari_grup_suzgeci():
    b = _bizim()
    p = Y.parti_satirlari(b, BUGUN, grup="180+ gün")
    assert p == [{"SKU": "A", "Ürün": "Ürün A", "Belge": "D1", "Tür": "İthalat", "Depoya giriş": "2026-03-08",
                  "Kalan adet": 20, "Yaş (gün)": 210, "Yaş grubu": "180+ gün"}]
    assert {x["Belge"] for x in Y.parti_satirlari(b, BUGUN)} == {"D1", "D2", "YI-1"}


def test_excel_sayfalari():
    import io
    import openpyxl
    b = _bizim()
    xb = Y.excel_bytes({"Ürünler": Y.urun_satirlari(b, {}, grup="180+ gün"), "Boş": []})
    wb = openpyxl.load_workbook(io.BytesIO(xb))
    assert wb.sheetnames == ["Ürünler", "Boş"]
    ws = wb["Ürünler"]
    basliklar = [c.value for c in ws[1]]
    assert "_id" not in basliklar and basliklar[:3] == ["SKU", "Ürün", "Kategori"]
    assert ws.cell(2, 1).value == "A" and wb["Boş"].cell(1, 1).value == "Bilgi"


# ── Stok kartı: stoğu biten ürün "sayım yüklenmemiş" demesin; ipucu takılmasın (Ekim 2026) ──
def test_stogu_biten_urunde_sayim_uyarisi_yok():
    from pathlib import Path
    k = (Path(__file__).resolve().parent.parent / "kayranpm" / "stok_karti.py").read_text(encoding="utf-8")
    i = k.index('elif alimlar or satislar or isinstance(urun.get("depo_kirilim"), dict):')
    j = k.index('pencere_bos("G5F depo sayımı yüklenmemiş')
    assert i < j and '_srow(d, 0, 0' in k[i:j] and "Bizim depolarda stok yok" not in k
    w = k[k.index("Başlangıç stoğu (Excel) yüklenmemiş") - 400:k.index("Başlangıç stoğu (Excel) yüklenmemiş")]
    assert "not (alimlar or satislar" in w


def test_kenar_arama_dugmesinde_ipucu_yok_ve_pencerede_ipucu_gizli():
    from pathlib import Path
    kok = Path(__file__).resolve().parent.parent
    m = (kok / "kayranpm" / "main.py").read_text(encoding="utf-8")
    d = m[m.index('key=f"stok_ac_{_r[\'sku\']}"') - 200:m.index('key=f"stok_ac_{_r[\'sku\']}"') + 120]
    assert "help=" not in d
    t = (kok / "shared" / "tasarim.py").read_text(encoding="utf-8")
    assert 'body:has(div[role="dialog"]) [data-baseweb="tooltip"]' in t


def test_satisi_eski_urunde_satis_verisi_yok_yazmaz():
    from pathlib import Path
    k = (Path(__file__).resolve().parent.parent / "kayranpm" / "stok_karti.py").read_text(encoding="utf-8")
    g = k[k.index("_son_satis = max("):k.index('_yeter = "satış verisi yok"')]
    assert "elif satislar:" in g and "stokta yok" in g and "son satış" in g
